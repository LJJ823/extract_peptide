#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
batch_extract_nanopore_signal_relations.py

核心目标：
    从主论文 + SI 中批量抽取纳米孔阻断信号关系，而不是优先补全多肽基本性质。

抽取对象：
    1. 阻断电流 / 残余电流类：
       Ib/I0, I/I0, residual current, relative current, normalized current,
       blockade current, current blockage, current drop, ΔI 等。
    2. 阻断时间类：
       tD, log(tD), dwell time, residence time, blockade time,
       translocation time, event duration 等。

重点建立关系：
    待测物 analyte
    纳米孔 nanopore
    阻断信号 metric
    图号 / panel
    电压 voltage
    待测物浓度 analyte concentration
    chamber / buffer / electrolyte / pH / temperature
    证据 evidence

说明：
    - 本脚本暂时不计算分子量、pI、净电荷、GRAVY 等性质。
    - σb、sigma_b、standard deviation、current fluctuation 等只作为辅助上下文，
      默认不作为目标信号输出。
    - 支持从图注中的 "+80 to +160 mV" 自动展开为 80/100/120/140/160 mV。
    - 支持可选 VLM 读 PDF 页面截图中的图、表、坐标轴、图例、图注。

依赖：
    pip install pymupdf pandas openpyxl requests
    # 如需读取 docx：
    pip install python-docx

环境变量：
    ARK_API_KEY / DOUBAO_API_KEY / VOLCENGINE_API_KEY
        豆包/火山方舟 API Key，三者任选其一。
    DOUBAO_MODEL / ARK_MODEL / VOLCENGINE_MODEL
        文本模型或推理接入点 ID。火山方舟上常见为 ep-... 接入点 ID；
        也可以按你的账号可用模型填写。
    DOUBAO_VL_MODEL / ARK_VL_MODEL
        可选，视觉模型或视觉推理接入点 ID，仅 --use-vlm-pages 时使用。
    ARK_BASE_URL
        可选，默认 https://ark.cn-beijing.volces.com/api/v3/chat/completions

用法：
    python batch_extract_nanopore_signal_relations.py "D:/Code/Python/extract_peptide"

启用 VLM 读图：
    python batch_extract_nanopore_signal_relations.py "D:/Code/Python/extract_peptide" --use-vlm-pages

限制 VLM 每个文件夹最多读多少页：
    python batch_extract_nanopore_signal_relations.py "D:/Code/Python/extract_peptide" --use-vlm-pages --vlm-max-pages-per-folder 20
"""

from __future__ import annotations

import argparse
import base64
import json
import math
import os
import re
import sys
import time
import traceback
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

import pandas as pd
import requests

try:
    import fitz  # PyMuPDF
except Exception as exc:
    raise RuntimeError("请先安装 PyMuPDF：pip install pymupdf") from exc

try:
    from docx import Document  # type: ignore
except Exception:
    Document = None


# =============================================================================
# 1. 基础配置
# =============================================================================

SUPPORTED_EXTS = {".pdf", ".docx", ".txt", ".text"}

SI_HINTS = [
    "si", "supp", "supplement", "supplementary", "supporting",
    "esm", "appendix", "附录", "补充"
]

SKIP_FOLDERS = {
    "__pycache__", ".git", ".idea", ".vscode", ".ipynb_checkpoints"
}

DEFAULT_BASE_URL = "https://ark.cn-beijing.volces.com/api/v3/chat/completions"

DOI_REGEX = re.compile(r"\b10\.\d{4,9}/[-._;()/:A-Z0-9]+\b", re.I)

UNICODE_SUBSCRIPT_MAP = str.maketrans("₀₁₂₃₄₅₆₇₈₉₊₋₍₎", "0123456789+-()")
UNICODE_SUPERSCRIPT_MAP = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻⁽⁾", "0123456789+-()")


# =============================================================================
# 2. 关键词和正则
# =============================================================================

SIGNAL_KEYWORDS: Dict[str, float] = {
    # nanopore
    "nanopore": 12,
    "pore": 4,
    "α-hemolysin": 10,
    "alpha-hemolysin": 10,
    "α-hl": 10,
    "alpha-hl": 10,
    "aerolysin": 12,
    "mspa": 10,
    "clya": 10,
    "frac": 10,
    "ompf": 8,
    "ompg": 8,
    "solid-state nanopore": 12,
    "solid state nanopore": 12,
    "silicon nitride": 8,
    "si3n4": 8,

    # current
    "blockade current": 18,
    "blockage current": 16,
    "current blockade": 16,
    "residual current": 18,
    "relative current": 14,
    "normalized current": 14,
    "current drop": 13,
    "current decrease": 10,
    "current blockage": 14,
    "event amplitude": 12,
    "blockade amplitude": 12,
    "open pore current": 8,
    "ionic current": 8,
    "ib/i0": 20,
    "i_b/i_0": 20,
    "i/i0": 18,
    "i / i0": 18,
    "δi": 12,
    "Δi": 12,
    "delta i": 10,

    # time
    "blockade time": 18,
    "blockage time": 16,
    "dwell time": 18,
    "residence time": 14,
    "translocation time": 14,
    "event duration": 14,
    "duration time": 10,
    "td": 8,
    "t_d": 10,

    # condition
    "voltage": 10,
    "bias": 8,
    "applied potential": 8,
    "transmembrane potential": 8,
    "mv": 5,
    "buffer": 6,
    "ph": 6,
    "kcl": 6,
    "nacl": 4,
    "licl": 4,
    "electrolyte": 6,
    "cis chamber": 8,
    "trans chamber": 8,
    "concentration": 6,

    # source
    "figure": 4,
    "fig.": 4,
    "table": 4,
    "caption": 5,
    "panel": 4,
}

SIGNAL_KEYWORDS_CN: Dict[str, float] = {
    "纳米孔": 12,
    "阻断电流": 18,
    "阻塞电流": 18,
    "残余电流": 18,
    "相对电流": 14,
    "归一化电流": 14,
    "电流下降": 12,
    "阻断幅度": 12,
    "阻断时间": 18,
    "阻塞时间": 16,
    "停留时间": 16,
    "驻留时间": 16,
    "易位时间": 14,
    "事件持续时间": 12,
    "电压": 10,
    "跨膜电压": 10,
    "缓冲液": 6,
    "电解质": 6,
    "图注": 5,
    "图": 4,
    "表": 4,
}

TARGET_CURRENT_RE = re.compile(
    r"(?:"
    r"I\s*_?b\s*/\s*I\s*_?0|I\s*/\s*I\s*0|Ib/I0|I/I0|I_b/I_0|"
    r"residual\s+current|relative\s+current|normalized\s+current|"
    r"blockade\s+current|blockage\s+current|current\s+blockade|"
    r"current\s+drop|current\s+decrease|current\s+blockage|"
    r"blockade\s+amplitude|event\s+amplitude|ΔI|δI|delta\s+I|"
    r"残余电流|相对电流|归一化电流|阻断电流|阻塞电流|电流下降|阻断幅度"
    r")",
    re.I,
)

TARGET_TIME_RE = re.compile(
    r"(?:"
    r"t\s*_?D|tD|log\s*\(?\s*t\s*_?D\s*\)?|"
    r"dwell\s+time|residence\s+time|blockade\s+time|blockage\s+time|"
    r"duration\s+time|event\s+duration|translocation\s+time|"
    r"阻断时间|阻塞时间|停留时间|驻留时间|易位时间|事件持续时间"
    r")",
    re.I,
)

NOISE_RE = re.compile(
    r"(?:σ\s*_?b|sigma\s*_?b|standard\s+deviation|std\.?|current\s+fluctuation|"
    r"blockade\s+noise|noise|电流波动|标准差)",
    re.I,
)

COUNTS_RE = re.compile(
    r"(?:^|\b)(counts?|number\s+of\s+events?|events?)(?:\b|$)|计数|事件数",
    re.I,
)

FIGURE_RE = re.compile(
    r"\b(?:Fig\.?|Figure|Table)\s*S?\d+[A-Za-z]?(?:[-–]\w+)?|图\s*S?\d+[A-Za-z]?|表\s*S?\d+[A-Za-z]?",
    re.I,
)

MV_SINGLE_RE = re.compile(r"([+-]?\d+(?:\.\d+)?)\s*mV\b", re.I)

STRICT_VOLTAGE_RANGE_RE = re.compile(
    r"(?:from|range(?:s|d)?\s*from|ranging\s*from|between|"
    r"applied\s+voltages?\s+(?:ranging\s+from)?|"
    r"voltage(?:s)?\s+(?:ranging\s+from|from)?|"
    r"bias(?:es)?\s+(?:ranging\s+from|from)?|"
    r"potential(?:s)?\s+(?:ranging\s+from|from)?)"
    r"[^\n.;]{0,60}?([+-]?\d+(?:\.\d+)?)\s*(?:to|[-–~]|and)\s*([+-]?\d+(?:\.\d+)?)\s*mV",
    re.I,
)

VOLTAGE_CONTEXT_RE = re.compile(
    r"(?:"
    r"Voltage\s*\(\s*mV\s*\)|applied\s+voltages?|voltages?|bias(?:\s+voltage)?|"
    r"transmembrane\s+(?:voltage|potential)|holding\s+potential|applied\s+potential|"
    r"x\s*[- ]?axis[^\n.;]{0,80}?(?:Voltage|mV)|电压|跨膜电压"
    r")",
    re.I,
)

MUTATION_OR_PEPTIDE_NUMBER_PATTERNS = [
    re.compile(r"(?:Aβ|A\s*β|Abeta|amyloid\s*beta|amyloid-?β)\s*\d+\s*[-–]\s*\d+", re.I),
    re.compile(r"\b[A-Z]\d{1,4}[A-Z]\b"),
    re.compile(r"\b[A-Z]\d{1,4}[A-Z](?:\s*/\s*[A-Z]\d{1,4}[A-Z])+\b"),
    re.compile(r"\b(?:peptide|residue(?:s)?|amino\s*acid(?:s)?|fragment)\s*\d+\s*[-–]\s*\d+\b", re.I),
]


# =============================================================================
# 3. 数据结构
# =============================================================================

@dataclass
class TextChunk:
    file_name: str
    file_type: str
    page_start: int
    page_end: int
    text: str
    score: float


# =============================================================================
# 4. 通用工具
# =============================================================================

def normalize_unicode(text: Any) -> str:
    if text is None:
        return ""
    return str(text).translate(UNICODE_SUBSCRIPT_MAP).translate(UNICODE_SUPERSCRIPT_MAP)


def normalize_whitespace(text: str) -> str:
    text = normalize_unicode(text)
    text = text.replace("\u00ad", "")
    text = text.replace("\xa0", " ")
    text = text.replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def clean_unknown(v: Any) -> Any:
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        if not s:
            return None
        if s.lower() in {"null", "none", "unknown", "not reported", "n/a", "na", "无", "未知"}:
            return None
        return s
    return v


def safe_float(x: Any) -> Optional[float]:
    x = clean_unknown(x)
    if x is None:
        return None
    if isinstance(x, (int, float)):
        if math.isnan(float(x)):
            return None
        return float(x)
    m = re.search(r"[-+]?\d+(?:\.\d+)?", str(x))
    return float(m.group(0)) if m else None


def json_dumps(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, indent=2)


def json_or_none(obj: Any) -> Optional[str]:
    if obj is None:
        return None
    if isinstance(obj, str):
        s = obj.strip()
        return s if s and s.lower() not in {"null", "none", "unknown", "n/a", "na"} else None
    try:
        return json.dumps(obj, ensure_ascii=False)
    except Exception:
        return str(obj)


def deduplicate_list(items: Iterable[Any]) -> List[Any]:
    out = []
    seen = set()
    for item in items:
        if item is None:
            continue
        key = str(item).strip()
        if not key or key.lower() in {"null", "none", "unknown", "n/a", "na"}:
            continue
        if key in seen:
            continue
        seen.add(key)
        out.append(item)
    return out


def extract_json_from_text(text: str) -> Any:
    text = (text or "").strip()

    try:
        return json.loads(text)
    except Exception:
        pass

    # 去掉 ```json ... ```
    text2 = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
    text2 = re.sub(r"\s*```$", "", text2)
    try:
        return json.loads(text2)
    except Exception:
        pass

    # 截取最外层 JSON
    starts = [i for i in [text.find("{"), text.find("[")] if i != -1]
    if not starts:
        # 打印原始输出以便调试
        print(f"[DEBUG] 模型输出中未找到 JSON 起始符号。原始输出前500字符：")
        print(text[:500])
        raise ValueError("模型输出中未找到 JSON 起始符号。")

    start = min(starts)
    for end in range(len(text), start, -1):
        snippet = text[start:end]
        try:
            return json.loads(snippet)
        except Exception:
            continue

    # 打印原始输出以便调试
    print(f"[DEBUG] 无法解析 JSON。原始输出前1000字符：")
    print(text[:1000])
    print(f"[DEBUG] 尝试解析的片段（最后500字符）：")
    print(text[-500:])
    raise ValueError("无法从模型输出中解析 JSON。")


# =============================================================================
# 5. Doubao / Volcano Ark LLM/VLM
# =============================================================================

class DoubaoArkLLM:
    def __init__(
        self,
        api_key: Optional[str] = None,
        model: Optional[str] = None,
        base_url: Optional[str] = None,
        timeout: int = 600,
        max_retries: int = 4,
    ) -> None:
        self.api_key = (
            api_key
            or os.getenv("ARK_API_KEY", "").strip()
            or os.getenv("DOUBAO_API_KEY", "").strip()
            or os.getenv("VOLCENGINE_API_KEY", "").strip()
        )
        self.model = (
            model
            or os.getenv("DOUBAO_MODEL", "").strip()
            or os.getenv("ARK_MODEL", "").strip()
            or os.getenv("VOLCENGINE_MODEL", "").strip()
            or "doubao-seed-2-0-pro-260215"
        )
        self.base_url = (
            base_url
            or os.getenv("ARK_BASE_URL", "").strip()
            or os.getenv("DOUBAO_BASE_URL", "").strip()
            or DEFAULT_BASE_URL
        )
        self.timeout = timeout
        self.max_retries = max_retries
        if not self.api_key:
            raise ValueError(
                "未检测到豆包/火山方舟 API Key。请先设置环境变量："
                "set ARK_API_KEY=你的key  或  set DOUBAO_API_KEY=你的key"
            )

    def chat(
        self,
        messages: List[Dict[str, Any]],
        temperature: float = 0.05,
        max_tokens: int = 5000,
    ) -> str:
        payload = {
            "model": self.model,
            "messages": messages,
            "temperature": temperature,
            "max_tokens": max_tokens,
        }
        headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json",
        }

        last_err: Optional[Exception] = None
        for attempt in range(1, self.max_retries + 1):
            try:
                resp = requests.post(
                    self.base_url,
                    headers=headers,
                    json=payload,
                    timeout=self.timeout,
                )
                if resp.status_code >= 400:
                    raise RuntimeError(f"HTTP {resp.status_code}: {resp.text[:1200]}")
                data = resp.json()
                return data["choices"][0]["message"]["content"]
            except Exception as exc:
                last_err = exc
                if attempt == self.max_retries:
                    break
                print(f"      [重试 {attempt}/{self.max_retries}] 模型调用失败: {exc}")
                time.sleep(min(20, 2 ** attempt))

        raise RuntimeError(f"大模型调用失败：{last_err}")


# =============================================================================
# 6. 文件读取
# =============================================================================

def read_pdf_pages(pdf_path: Path) -> List[Tuple[int, str]]:
    pages: List[Tuple[int, str]] = []
    doc = fitz.open(pdf_path)
    try:
        for i in range(doc.page_count):
            text = doc.load_page(i).get_text("text") or ""
            text = normalize_whitespace(text)
            if text:
                pages.append((i + 1, text))
    finally:
        doc.close()
    return pages


def read_docx_text(docx_path: Path) -> List[Tuple[int, str]]:
    if Document is None:
        print(f"  - {docx_path.name}: 未安装 python-docx，跳过 docx")
        return []
    doc = Document(str(docx_path))
    paras = [normalize_whitespace(p.text) for p in doc.paragraphs if p.text and p.text.strip()]
    text = "\n".join(paras)
    return [(1, text)] if text.strip() else []


def read_txt_text(txt_path: Path) -> List[Tuple[int, str]]:
    text = txt_path.read_text(encoding="utf-8", errors="ignore")
    text = normalize_whitespace(text)
    return [(1, text)] if text.strip() else []


def read_file_pages(file_path: Path) -> List[Tuple[int, str]]:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf_pages(file_path)
    if suffix == ".docx":
        return read_docx_text(file_path)
    if suffix in {".txt", ".text"}:
        return read_txt_text(file_path)
    return []


def classify_file_role(file_path: Path) -> str:
    lower = file_path.name.lower()
    if any(h in lower for h in SI_HINTS):
        return "SI"
    return "main_or_other"


def gather_supported_files(folder: Path) -> List[Path]:
    return sorted(
        p for p in folder.iterdir()
        if p.is_file() and p.suffix.lower() in SUPPORTED_EXTS
    )


# =============================================================================
# 7. Chunk 与候选筛选
# =============================================================================

def score_chunk(text: str, file_name: str) -> float:
    raw = normalize_unicode(text or "")
    lower = raw.lower()
    score = 0.0

    for kw, w in SIGNAL_KEYWORDS.items():
        if kw in lower:
            score += w

    for kw, w in SIGNAL_KEYWORDS_CN.items():
        if kw in raw:
            score += w

    if TARGET_CURRENT_RE.search(raw):
        score += 25
    if TARGET_TIME_RE.search(raw):
        score += 25
    if FIGURE_RE.search(raw):
        score += 8
    if MV_SINGLE_RE.search(raw):
        score += 8
    if re.search(r"\b(?:cis|trans)\s+chamber\b", raw, re.I):
        score += 8
    if re.search(r"\d+(?:\.\d+)?\s*(?:μM|uM|µM|mM|nM|pM)\b", raw, re.I):
        score += 6

    # 文件名里 SI 略微加权
    if any(h in file_name.lower() for h in SI_HINTS):
        score += 1.5

    return score


def chunk_pages(
    file_name: str,
    file_type: str,
    pages: List[Tuple[int, str]],
    chunk_size: int = 9000,
    overlap: int = 800,
) -> List[TextChunk]:
    chunks: List[TextChunk] = []
    buffer = ""
    start_page: Optional[int] = None
    last_page: Optional[int] = None

    def flush(buf: str, p1: int, p2: int) -> None:
        text = normalize_whitespace(buf)
        if not text:
            return
        chunks.append(
            TextChunk(
                file_name=file_name,
                file_type=file_type,
                page_start=p1,
                page_end=p2,
                text=text,
                score=score_chunk(text, file_name),
            )
        )

    for page_no, page_text in pages:
        page_blob = f"\n[Page {page_no}]\n{page_text}\n"
        if start_page is None:
            start_page = page_no

        if len(buffer) + len(page_blob) <= chunk_size:
            buffer += page_blob
            last_page = page_no
            continue

        if buffer:
            flush(buffer, start_page or page_no, last_page or page_no)
            carry = buffer[-overlap:] if overlap > 0 else ""
            buffer = carry + page_blob
            start_page = page_no if not carry else max(1, page_no - 1)
            last_page = page_no
        else:
            flush(page_blob, page_no, page_no)
            buffer = ""
            start_page = None
            last_page = None

    if buffer and start_page is not None and last_page is not None:
        flush(buffer, start_page, last_page)

    return chunks


def select_signal_candidate_chunks(
    chunks: List[TextChunk],
    top_k: int = 45,
    min_score: float = 10.0,
) -> List[TextChunk]:
    hits = [c for c in chunks if c.score >= min_score]
    return sorted(hits or chunks, key=lambda c: c.score, reverse=True)[:top_k]


# =============================================================================
# 8. VLM 读 PDF 页面截图
# =============================================================================

def render_pdf_page_to_base64(pdf_path: Path, page_no_1based: int, zoom: float = 2.0) -> str:
    doc = fitz.open(pdf_path)
    try:
        page = doc.load_page(page_no_1based - 1)
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom), alpha=False)
        return base64.b64encode(pix.tobytes("png")).decode("ascii")
    finally:
        doc.close()


def choose_vlm_pages(
    files: List[Path],
    pages_map: Dict[str, List[Tuple[int, str]]],
    signal_chunks: List[TextChunk],
    max_pages: int,
) -> List[Tuple[Path, int]]:
    scores: Dict[Tuple[str, int], float] = defaultdict(float)

    for ch in signal_chunks:
        for p in range(ch.page_start, ch.page_end + 1):
            scores[(ch.file_name, p)] += ch.score

    for fp in files:
        if fp.suffix.lower() != ".pdf":
            continue
        for p, text in pages_map.get(fp.name, []):
            page_score = score_chunk(text, fp.name)
            if FIGURE_RE.search(text):
                page_score += 12
            if re.search(r"\b(?:Fig\.?|Figure)\s*S?\d+", text, re.I):
                page_score += 8
            scores[(fp.name, p)] += min(page_score, 60)

    name_to_path = {p.name: p for p in files}
    chosen: List[Tuple[Path, int]] = []
    for (fname, page_no), _score in sorted(scores.items(), key=lambda x: x[1], reverse=True):
        fp = name_to_path.get(fname)
        if fp and fp.suffix.lower() == ".pdf":
            chosen.append((fp, page_no))
        if len(chosen) >= max_pages:
            break

    return chosen


def vlm_extract_page_signal_text(vlm: DoubaoArkLLM, pdf_path: Path, page_no: int) -> Optional[TextChunk]:
    try:
        b64 = render_pdf_page_to_base64(pdf_path, page_no)
    except Exception as exc:
        print(f"    VLM 渲染失败: {pdf_path.name} page {page_no}: {exc}")
        return None

    prompt = (
        "请识别这页论文截图中的图、表、图注、坐标轴、图例和面板编号。"
        "只提取纳米孔目标阻断信号关系："
        "1) Ib/I0、I/I0、residual current、blockade current、current drop 等阻断电流/残余电流类；"
        "2) tD、log(tD)、dwell time、blockade time、residence time 等阻断/停留时间类。"
        "σb、sigma_b、standard deviation、current fluctuation 是电流波动/标准差，不是本任务目标，"
        "不要把 σb 面板单独作为目标记录，只可作为 paired_panels 或 evidence 上下文。"
        "必须逐个 panel 读取 x_axis_label、x_axis_unit、y_axis_label、y_axis_unit、chart_type。"
        "极其重要：必须判断每个 panel 展示的是电流指标还是时间指标，用两个字段记录："
        "- 如果 panel a 展示 Ib/I0、residual current 等电流指标 → panel_dwell_current='a'"
        "- 如果 panel b 展示 tD、dwell time 等时间指标 → panel_dwell_time='b'"
        "- 非目标 panel（如 σb）不填入任何字段"
        "面板维度优先级：当同一张图中存在多种维度的面板时，优先使用一维图（直方图、柱状图、电压依赖折线图等）。"
        "如果一维图已经能读到 dwell_current 和 dwell_time 的数值，就不再考虑二维图乃至三维图。"
        "三维散点图、等高线图、热力图永远不作为目标数据来源。"
        "目标指标可能在 y 轴，也可能在 x 轴：如果 x 轴是 Ib/I0 或 log(tD)/tD，y 轴是 Counts/事件数，也要保留。"
        "如果 x 轴是 Voltage (mV)，例如 80、100、120、140、160，请完整列出 voltage_series_mV，不能只列一个中间值。"
        "极其重要：不要把待测物名称、肽段编号或突变编号中的数字当成电压，例如 Aβ18-26 中的 18 和 26、"
        "A21G 中的 21、T232K/K238Q 中的 232 和 238 都不是电压。"
        "只有图横坐标 Voltage(mV)、图注 applied voltages/ranging from ... mV 或正文明确 voltage/bias 的数字才是电压。"
        "请结合图注抽取纳米孔名称和突变体、待测物、加样侧、终浓度、buffer、电解质、pH、温度。"
        "如果图中有多个待测物，每个待测物都要单独列出。能估读目标数值就给出 values_by_voltage；读不准填 null 并说明 approximate/uncertain。"
        "没有相关目标信息则只输出 NO_RELEVANT_SIGNAL。"
    )

    messages: List[Dict[str, Any]] = [
        {
            "role": "user",
            "content": [
                {"type": "text", "text": prompt},
                {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{b64}"}},
            ],
        }
    ]

    try:
        raw = vlm.chat(messages, temperature=0.03, max_tokens=3500)
    except Exception as exc:
        print(f"    VLM 识别失败: {pdf_path.name} page {page_no}: {exc}")
        return None

    if not raw or "NO_RELEVANT_SIGNAL" in raw:
        return None

    text = f"[VLM_PAGE_IMAGE] file={pdf_path.name}; page={page_no}; source_kind=figure_image\n{raw}"
    return TextChunk(
        file_name=pdf_path.name,
        file_type="figure_image_vlm",
        page_start=page_no,
        page_end=page_no,
        text=text,
        score=score_chunk(text, pdf_path.name) + 35,
    )


def extract_vlm_chunks(
    files: List[Path],
    pages_map: Dict[str, List[Tuple[int, str]]],
    signal_chunks: List[TextChunk],
    vlm: DoubaoArkLLM,
    max_pages: int,
) -> List[TextChunk]:
    pages = choose_vlm_pages(files, pages_map, signal_chunks, max_pages)
    if not pages:
        return []

    print(f"  VLM 将识别候选图表页: {len(pages)}")
    out: List[TextChunk] = []
    for fp, page_no in pages:
        print(f"    - VLM: {fp.name} page {page_no}")
        ch = vlm_extract_page_signal_text(vlm, fp, page_no)
        if ch:
            out.append(ch)
    return out


# =============================================================================
# 9. LLM Prompt：阻断信号关系抽取
# =============================================================================

def build_signal_relation_prompt(folder_name: str, chunks: List[TextChunk]) -> List[Dict[str, Any]]:
    parts: List[str] = []
    total = 0
    max_chars = 80000

    for i, ch in enumerate(chunks, 1):
        block = (
            f"\n===== CHUNK {i} =====\n"
            f"file_name: {ch.file_name}\n"
            f"file_type: {ch.file_type}\n"
            f"pages: {ch.page_start}-{ch.page_end}\n"
            f"score: {ch.score:.2f}\n"
            f"text:\n{ch.text}\n"
        )
        if total + len(block) > max_chars:
            break
        parts.append(block)
        total += len(block)

    context = "\n".join(parts)

    system = (
        "你是纳米孔单分子检测论文的信息抽取助手。"
        "你的核心任务是建立阻断电流/阻断时间与待测物、纳米孔、实验条件、浓度、图表来源之间的对应关系。"
        "注意：一篇论文可能包含多种不同的纳米孔（不同孔蛋白、不同突变体、生物孔 vs 固态孔），每种纳米孔的数据必须分开记录。"
        "当前任务不需要计算待测物分子量、pI、净电荷、GRAVY 等基本性质。"
        "必须输出严格 JSON，不要输出解释性文字。"
    )

    user = f"""
请从子文件夹 {folder_name} 的全文文字、表格、图注、正文引用图的位置、以及 VLM 识别到的图中文字中，抽取所有纳米孔阻断信号关系。

你需要抽取的目标信号只有两大类：

A. 阻断电流 / 残余电流类：
   Ib/I0, I/I0, I_b/I_0, residual current, relative current,
   normalized current, blockade current, current drop, current blockage,
   ΔI, δI, blockade amplitude 等。

B. 阻断时间类：
   tD, t_D, log(tD), dwell time, residence time,
   blockade time, blockage time, translocation time, event duration 等。

不要把以下内容作为目标信号输出：
   σb, sigma_b, standard deviation, std, current fluctuation, blockade noise。
这些可以写在 paired_panels、panel_relation 或 evidence 中作为辅助上下文，但不要单独输出为目标记录。

关键规则：
1. 核心关系链：待测物 analyte ↔ 纳米孔 nanopore ↔ 实验条件 ↔ 图/面板 ↔ 信号指标。
   这五项是一个不可拆分的整体。同一待测物在不同纳米孔中测得的数据属于不同记录，同一纳米孔测不同待测物可以共享 nanopore 字段但 analyte 必须分开。
2. 纳米孔和待测物必须从同一上下文中配对抽取，不能各自独立抽取后随意组合。
   典型的上下文来源是图注（figure caption），例如：
     “Detection of the WT and A21G Aβ18-26 peptides with the WT aerolysin nanopore.”
   从这句话应得到：
     - nanopore_name = “aerolysin nanopore”, nanopore_variant = “WT”
     - analytes = [“WT Aβ18-26”, “A21G Aβ18-26”]
     - 两个待测物共享同一个纳米孔，可各自生成记录但 nanopore 字段相同。
   再如：
     “Figure 3 shows α-hemolysin measurements; Figure 5 uses aerolysin for the same peptide.”
   则 Figure 3 的记录 nanopore = α-hemolysin，Figure 5 的 nanopore = aerolysin，绝不能混淆。
3. **待测物必须完整列出，严禁遗漏**（极其重要）：
   3a. 一个 Figure 中经常同时比较多个待测物。例如 Figure S6 同时比较 WT Aβ18-26 和 A21G Aβ18-26，
       必须两个都输出，不能只输出其中一个。
   3b. **图注中列出的所有待测物，一个都不能少**：
       仔细阅读图注全文，找出所有被提及的待测物名称。图注中常见的多待测物描述模式：
       - "for the WT and A21G Aβ18-26 peptides" → 2 个待测物
       - "comparison of peptide A, peptide B, and peptide C" → 3 个待测物
       - "various concentrations of..." → 可能有多个浓度条件
   3c. **图中出现的所有待测物，一个都不能少**：
       图中不同曲线/不同颜色/不同形状的标记通常对应不同待测物。必须把图中每个图例项（legend entry）
       对应的待测物都识别出来。即使图注没有在文字中逐一列举，图例中出现的每个待测物也必须输出。
   3d. **交叉验证**：提取完成后，必须自查——图注中提到了 N 个待测物名称，图中图例显示了 M 个待测物，
       最终输出的记录中应覆盖这 N 和 M 的并集，不得遗漏任何一个。
   3e. 如果图中使用了不同颜色/形状区分待测物（如黑色方形 vs 红色圆形），必须根据图例逐一识别，
       不要把不同待测物的数据混在一起。
4. 纳米孔突变体必须保留。例如 T232K/K238Q aerolysin nanopore。
4b. 一篇论文里可能同时使用了多种不同的纳米孔（例如同时测了 α-hemolysin 和 aerolysin，或者同时测了生物孔和固态孔）。
    每种纳米孔必须作为独立维度分别输出，绝不能把不同纳米孔的数据混在同一条记录里。
    又如同一张图的 panel a 用 WT α-HL，panel b 用 M113F α-HL 突变体，也必须分别标注 nanopore_name + nanopore_variant。
5. 实验条件必须尽量抽取：
   voltage、buffer、pH、electrolyte、salt_concentration、temperature、cis/trans chamber、analyte_concentration。
   5a. **待测物浓度必须按 analyte 区分，不同浓度绝不能合并**（极其重要）：
       如果图注或正文中明确描述了不同待测物有不同的浓度，必须分别记录，绝不能合并为一个笼统的浓度。
       例如图注写 "WT Aβ18-26 was added at a final concentration of 10.0 μM, and A21G Aβ18-26 at 5.0 μM"：
       - WT Aβ18-26 的记录：analyte_concentration_value=10.0, analyte_concentration_unit="μM"
       - A21G Aβ18-26 的记录：analyte_concentration_value=5.0, analyte_concentration_unit="μM"
       绝不能把两个浓度合并为 "10.0 and 5.0 μM"。
       如果图注只给了一个笼统浓度（如 "peptides were added at 10 μM"），且无法确定每个待测物的具体浓度，
       则所有待测物可以共享同一个浓度值，但必须在 info_source_notes 中注明浓度是笼统值。
       当 concentration_scope 为 "analyte-specific" 时，必须为不同浓度的待测物生成各自独立的记录。
6. 电压是关系键。不同电压不能合并。如果图或图注显示 +80 到 +160 mV，且图中电压点是 80、100、120、140、160，
   必须输出 voltage_series_mV=[80,100,120,140,160]。
7. 不能把待测物名称、肽段编号或突变编号中的数字当成电压：
   - Aβ18-26 中的 18 和 26 不是电压；
   - A21G 中的 21 不是电压；
   - T232K/K238Q 中的 232 和 238 不是电压。
8. 只有以下来源的数字才能作为电压：
   - 图横坐标 Voltage (mV) 的 tick；
   - 图注中的 applied voltages / ranging from ... mV；
   - 正文明确写的 voltage / bias / transmembrane potential / holding potential。
9. **panel 拆分为 panel_dwell_time 和 panel_dwell_current 两列**（极其重要）：
   9a. **核心思路**：同一个 figure 的不同 panel 可能展示不同类型的指标。
       例如 Figure S6 的 caption：
         "Voltage dependencies of Ib/I0 (a), tD (b), and σb (c) for the WT and A21G Aβ18-26
          peptides with the applied voltages ranging from +80 to +160 mV."
       其中 panel a 展示电流指标 Ib/I0，panel b 展示时间指标 tD。
       你应该判断每个 panel 对应哪种指标类型，然后用两个字段记录：
       - panel_dwell_current = "a" （panel a 对应电流指标）
       - panel_dwell_time = "b" （panel b 对应时间指标）
   9b. **每个待测物只输出一条记录**，不要为每个 panel 单独输出一条记录。
       上例中 WT Aβ18-26 应该输出一条记录：
       {{
         "figure_id": "Figure S6",
         "panel_dwell_current": "a",
         "panel_dwell_time": "b",
         "analytes": ["WT Aβ18-26", "A21G Aβ18-26"],
         "dwell_current_metric": "Ib/I0",
         "dwell_time_metric": "tD",
         ...
       }}
       不要输出 4 条记录（每个 panel × 每个待测物 各一条）。
   9c. **panel 字段的填写规则**：
       - 如果某个 panel 展示的是电流指标（Ib/I0、residual current 等）→ 填入 panel_dwell_current
       - 如果某个 panel 展示的是时间指标（tD、dwell time 等）→ 填入 panel_dwell_time
       - 如果某个 panel 展示的是非目标指标（如 σb）→ 不填，对应字段留 null
       - 如果只有一个 panel 且是电流指标 → panel_dwell_current="a", panel_dwell_time=null
       - 如果只有一个 panel 且是时间指标 → panel_dwell_time="a", panel_dwell_current=null
   9d. **dwell_current_metric 和 dwell_time_metric**（必填）：
       - dwell_current_metric: 电流指标的具体名称，如 "Ib/I0"、"residual current" 等
       - dwell_time_metric: 时间指标的具体名称，如 "tD"、"dwell time" 等
       - 如果某类指标不存在，对应字段留 null
   9e. **values_by_voltage 的合并**：
       - dwell_current_values_by_voltage: 电流指标按电压的值列表
       - dwell_time_values_by_voltage: 时间指标按电压的值列表
       - 两类指标的值分开存放，不要混在一起
   9f. 如果有多个待测物，每个待测物一条记录，但 panel 映射关系相同。
       例如 Figure S6 有 WT 和 A21G 两个待测物：
       - 记录1: analyte="WT Aβ18-26", panel_dwell_current="a", panel_dwell_time="b"
       - 记录2: analyte="A21G Aβ18-26", panel_dwell_current="a", panel_dwell_time="b"
   9g. **figure_id 拆分**（极其重要）：
       dwell_current 和 dwell_time 可能来自不同的 figure。例如：
       - Figure 4 的 panel b 展示 Ib/I0（电流指标），Figure 5 的 panel b 展示 tD（时间指标）
       此时必须为每种指标分别记录其来源 figure：
       - figure_id_dwell_current = "Figure 4"（电流指标来源）
       - figure_id_dwell_time = "Figure 5"（时间指标来源）
       如果两种指标来自同一个 figure，两个字段填相同的值。
       如果某种指标不存在，对应字段留 null。
10. 目标指标可能在 y 轴，也可能在 x 轴：
    - y 轴是 Ib/I0 或 tD(ms)：axis_role="target_on_y_axis"；
    - x 轴是 Ib/I0 或 log(tD)/tD，y 轴是 Counts：axis_role="target_on_x_axis_counts_on_y_axis"。
11. 如果图里能读出每个电压对应的目标信号值，可输出 values_by_voltage。
    读不准时 value 填 null，note 写 approximate/uncertain。
12. 不要臆造。无法确定填 null。
13. **信息来源优先级与上下文证明**（极其重要）：
    关键信息不需要必须出现在图注中，但必须能从该图对应的完整上下文中被可靠证明。
    大模型需要按以下优先级查找信息来源：
    优先级1：图注（figure caption）
    优先级2：正文中引用该图的位置（如 "As shown in Figure 3..."、"Figure 5 demonstrates..."）
    优先级3：图附近的上下文段落（图前后的描述性段落）
    优先级4：方法部分 / 全局实验条件（如 "All experiments were performed in 1 M KCl, pH 8.0"）
    只有当图对应的完整上下文中可以明确确定实验条件时才提取。上下文包括：
    - 图注（figure caption）
    - 正文中引用该图的段落
    - 图前后相关段落
    - 方法部分中的默认实验条件
    如果某个信息没有出现在图注中，但在正文或方法部分中被明确说明，并且可以合理对应到该图，
    则允许提取，并在 info_source_notes 字段中标明该信息来源（如 "voltage from methods section"、
    "buffer from in-text reference to Figure 3" 等）。
    如果无法确定该信息是否适用于该图，则不要提取该信息。
    对于每条记录，请输出 info_source_notes 字段，说明各关键信息（nanopore、analyte、metric、
    voltage、buffer、pH 等）的具体来源。
14. **图注和正文中的逐 panel / 逐 Figure 描述**（极其重要）：
    14a. 图注（figure caption）中经常会对每个 panel 逐一描述，例如：
        "Figure 4. (a) Ib/I0 histogram for WT Aβ18-26 at 10 μM.
         (b) tD histogram for WT Aβ18-26 at 10 μM.
         (c) Ib/I0 histogram for A21G Aβ18-26 at 5.0 μM.
         (d) tD histogram for A21G Aβ18-26 at 5.0 μM."
        你必须逐 panel 读取其中的待测物、浓度、指标类型等信息，不能只看图注的第一句话。
        每个 panel 可能对应不同的待测物、不同的浓度、不同的实验条件，必须分别准确记录。
    14b. 正文中引用 Figure 的位置也可能对各个 panel 有单独描述，例如：
        "As shown in Figure 4a, the Ib/I0 of WT Aβ18-26 was 0.45 at 100 mV,
         while Figure 4b shows that its dwell time was 3.2 ms."
        你必须检查正文引用处是否有对特定 panel 的补充描述（如具体数值、条件说明）。
        如果正文引用处提供了图注中未包含的信息，必须提取并标注来源。
    14c. 当同一 Figure 的不同 panel 使用不同实验条件（如不同浓度、不同 buffer、不同待测物）时，
        必须准确识别每个 panel 对应的条件，并在对应记录中分别填写。
        不要将 panel a 的条件错误地应用到 panel b。
    14d. **逐 panel 核查待测物完整性**：
        图注中对各 panel 的逐条描述是待测物的权威来源。例如图注写：
        "(a) Ib/I0 for WT and A21G. (b) tD for WT and A21G. (c) Ib/I0 for A21G mutant alone."
        panel a 和 b 涉及 2 个待测物（WT + A21G），panel c 涉及 1 个待测物（A21G mutant alone）。
        你必须根据每个 panel 的描述分别确定该 panel 中出现了哪些待测物，确保全部覆盖。
        如果 VLM 读取到图中某 panel 有 3 条曲线（图例有 3 个待测物），但图注只提到 2 个，
        以图中实际出现的为准，三者全部输出。
        最终必须自查：图中每个 panel 的图例项数量 == 该 panel 输出记录中涉及的待测物数量。
请输出严格 JSON，格式如下：
{{
  "title": null,
  "doi": null,
  "records": [
    {{
      "record_id": "R1",
      "figure_id_dwell_current": "Figure S6",
      "figure_id_dwell_time": "Figure S6",
      "figure_group_id": "Figure S6",

      "panel_dwell_current": "a",
      "panel_dwell_time": "b",
      "panel_relation": "panel a=Ib/I0 (current) and panel b=tD (time) are target outputs; panel c=σb is non-target context",

      "analytes": ["WT Aβ18-26", "A21G Aβ18-26"],
      "analyte_name": null,
      "analyte_alias": [],
      "analyte_type": "peptide/protein/amino acid/compound/unknown",

      "nanopore_name": "aerolysin nanopore",
      "nanopore_variant": "T232K/K238Q",
      "nanopore_type": "biological/solid-state/unknown",

      "dwell_current_metric": "Ib/I0",
      "dwell_current_metric_category": "residual_current/blockade_current/unknown",
      "dwell_current_chart_type": "voltage_dependency/histogram/density_plot/scatter/trace/table/text/other",
      "dwell_current_axis_role": "target_on_y_axis/target_on_x_axis_counts_on_y_axis/target_on_x_axis/unknown",
      "dwell_current_x_axis_label": "Voltage (mV)",
      "dwell_current_x_axis_unit": "mV",
      "dwell_current_y_axis_label": "Ib/I0",
      "dwell_current_y_axis_unit": "dimensionless",

      "dwell_time_metric": "tD",
      "dwell_time_metric_category": "dwell_time/blockade_time/unknown",
      "dwell_time_chart_type": "voltage_dependency/histogram/density_plot/scatter/trace/table/text/other",
      "dwell_time_axis_role": "target_on_y_axis/target_on_x_axis_counts_on_y_axis/target_on_x_axis/unknown",
      "dwell_time_x_axis_label": "Voltage (mV)",
      "dwell_time_x_axis_unit": "mV",
      "dwell_time_y_axis_label": "tD",
      "dwell_time_y_axis_unit": "ms",

      "voltage_raw": "+80 to +160 mV",
      "voltage_mV": null,
      "voltage_series_raw": "+80, +100, +120, +140, +160 mV",
      "voltage_series_mV": [80, 100, 120, 140, 160],
      "voltage_polarity": "+",

      "dwell_current_values_by_voltage": [
        {{"voltage_mV": 80, "value": null, "unit": "Ib/I0", "note": "value not read from caption"}}
      ],
      "dwell_time_values_by_voltage": [
        {{"voltage_mV": 80, "value": null, "unit": "ms", "note": "value not read from caption"}}
      ],
      "dwell_current_signal_value": null,
      "dwell_current_signal_unit": null,
      "dwell_time_signal_value": null,
      "dwell_time_signal_unit": null,

      "analyte_concentration_raw": "final concentration of 10.0 μM for WT; 5.0 μM for A21G",
      "analyte_concentration_value": 10.0,
      "analyte_concentration_unit": "μM",
      "concentration_scope": "analyte-specific（每个待测物有独立浓度时必须标注）/ figure-wide / panel-specific / unknown",

      "chamber": "cis",
      "cis_trans_addition": "added to the cis chamber",
      "electrolyte": "1 M KCl",
      "salt_concentration": "1 M",
      "buffer": "10 mM Tris buffer",
      "pH": 8.0,
      "temperature": null,

      "source_file": null,
      "page_range": null,
      "source_kind": "text/table/figure_caption/figure_image/in_text_reference/unknown",
      "caption_text": "逐 panel 读取的图注原文，如 '(a) Ib/I0 histogram... (b) tD histogram...'",
      "in_text_reference_evidence": "正文引用该图时的逐 panel 描述，如 'Figure 4a shows Ib/I0 of WT...'，若有",
      "evidence": "最关键证据",
      "info_source_notes": {{
        "nanopore": "来源：图注/正文引用/方法部分",
        "analyte": "来源：图注",
        "metric": "来源：图注",
        "voltage": "来源：图注/正文引用/方法部分",
        "buffer": "来源：方法部分",
        "pH": "来源：方法部分",
        "concentration": "来源：图注"
      }},
      "ambiguous_reason": null,
      "confidence": 0.0
    }}
  ]
}}

输出前必须完成以下自查（极其重要）：
1. 逐 panel 检查：图注中每个 panel 的描述是否都已覆盖？有没有 panel 被跳过？
2. 逐待测物检查：图注中提到的所有待测物 + 图中图例出现的所有待测物 是否都已输出？
   总数是否匹配？有没有待测物被遗漏？
3. 逐浓度检查：不同待测物的浓度是否已分别记录？有没有被错误合并？
4. 如果自查发现遗漏，必须补充后再输出最终 JSON。

待抽取内容如下：
{context}
""".strip()

    return [
        {"role": "system", "content": system},
        {"role": "user", "content": user},
    ]


def llm_extract_signal_relations(llm: DoubaoArkLLM, folder_name: str, chunks: List[TextChunk]) -> Dict[str, Any]:
    raw = llm.chat(
        build_signal_relation_prompt(folder_name, chunks),
        temperature=0.03,
        max_tokens=30000,
    )
    # 调试：打印 LLM 返回的原始内容
    print(f"[DEBUG] LLM 返回内容长度: {len(raw)} 字符")
    print(f"[DEBUG] LLM 返回内容前500字符:")
    print(raw[:500])
    print(f"[DEBUG] ...")
    data = extract_json_from_text(raw)
    if not isinstance(data, dict):
        raise ValueError("LLM 返回根对象不是 dict")
    if not isinstance(data.get("records"), list):
        data["records"] = []
    data["_raw_response"] = raw
    return data


# =============================================================================
# 10. 标题/DOI
# =============================================================================

def coerce_title(title: Any, first_pages: List[Tuple[int, str]]) -> Optional[str]:
    title = clean_unknown(title)
    if isinstance(title, str):
        return title[:500]

    if first_pages:
        first_text = first_pages[0][1]
        lines = [x.strip() for x in first_text.splitlines() if x.strip()]
        # 尽量取前几个像标题的长行
        candidates = [ln for ln in lines[:20] if 20 <= len(ln) <= 300]
        if candidates:
            return candidates[0]
        if lines:
            return lines[0][:300]
    return None


def coerce_doi(doi: Any, all_texts: List[str]) -> Optional[str]:
    doi = clean_unknown(doi)
    if isinstance(doi, str) and DOI_REGEX.search(doi):
        return DOI_REGEX.search(doi).group(0)  # type: ignore[union-attr]
    for text in all_texts:
        m = DOI_REGEX.search(text)
        if m:
            return m.group(0)
    return None


# =============================================================================
# 11. 记录标准化：电压、浓度、目标信号、展开
# =============================================================================

def extract_numbers_from_mutation_or_peptide_names(*texts: Any) -> set:
    false_nums = set()
    joined = " | ".join(str(t or "") for t in texts)
    for pat in MUTATION_OR_PEPTIDE_NUMBER_PATTERNS:
        for m in pat.finditer(joined):
            for n in re.findall(r"\d+(?:\.\d+)?", m.group(0)):
                try:
                    false_nums.add(round(abs(float(n)), 6))
                except Exception:
                    pass
    return false_nums


def has_reliable_voltage_context(text: Any) -> bool:
    s = normalize_unicode(text)
    if not s.strip():
        return False
    if VOLTAGE_CONTEXT_RE.search(s):
        return True
    mv_hits = re.findall(r"[+-]?\d+(?:\.\d+)?\s*mV\b", s, flags=re.I)
    if len(mv_hits) >= 2 and re.search(r"voltage|bias|potential|mV|电压", s, flags=re.I):
        return True
    return False


def num_to_float(x: Any) -> Optional[float]:
    return safe_float(x)


def expand_range_values(a: float, b: float, text: str) -> List[float]:
    lo, hi = sorted([abs(float(a)), abs(float(b))])
    diff = hi - lo
    if diff <= 0:
        return [lo]

    step = None
    m_step = re.search(r"(?:step|interval|increment)[^\n.;]{0,50}?(\d+(?:\.\d+)?)\s*mV", text, flags=re.I)
    if m_step:
        step = float(m_step.group(1))

    # 常见电压依赖图：80,100,120,140,160
    if step is None:
        if diff <= 300 and abs((diff / 20) - round(diff / 20)) < 1e-6:
            step = 20.0
        elif diff <= 300 and abs((diff / 10) - round(diff / 10)) < 1e-6:
            step = 10.0

    if step and step > 0:
        vals = []
        cur = lo
        while cur <= hi + 1e-9 and len(vals) <= 100:
            vals.append(round(cur, 6))
            cur += step
        if abs(vals[-1] - hi) > 1e-6:
            vals.append(hi)
        return vals

    return [lo, hi]


def parse_voltage_series_from_text(
    text: Any,
    false_nums: Optional[set] = None,
    require_context: bool = True,
) -> List[float]:
    if text is None:
        return []
    false_nums = false_nums or set()
    s = normalize_unicode(text)
    if not s.strip():
        return []

    values: List[float] = []

    # 1. 严格范围：applied voltages ranging from +80 to +160 mV
    for m in STRICT_VOLTAGE_RANGE_RE.finditer(s):
        a = num_to_float(m.group(1))
        b = num_to_float(m.group(2))
        if a is not None and b is not None:
            values.extend(expand_range_values(a, b, s))

    # 2. 在 Voltage/Bias 上下文短语内读 tick 数字
    ctx_patterns = [
        r"(?:Voltage\s*\(\s*mV\s*\)|applied\s+voltages?|voltages?|bias(?:\s+voltage)?|transmembrane\s+(?:voltage|potential)|holding\s+potential|applied\s+potential|电压|跨膜电压)[^\n.;]{0,260}",
        r"[^\n.;]{0,100}(?:Voltage\s*\(\s*mV\s*\))[^\n.;]{0,260}",
    ]
    for pat in ctx_patterns:
        for m in re.finditer(pat, s, flags=re.I):
            phrase = m.group(0)

            for rm in STRICT_VOLTAGE_RANGE_RE.finditer(phrase):
                a = num_to_float(rm.group(1))
                b = num_to_float(rm.group(2))
                if a is not None and b is not None:
                    values.extend(expand_range_values(a, b, phrase))

            for n in re.findall(r"[+-]?\d+(?:\.\d+)?", phrase):
                v = num_to_float(n)
                if v is not None and 10 <= abs(v) <= 1000:
                    values.append(abs(v))

    # 3. 只有有可靠上下文时，才接受单个 mV
    if (not require_context) or has_reliable_voltage_context(s):
        for m in MV_SINGLE_RE.finditer(s):
            v = num_to_float(m.group(1))
            if v is not None:
                values.append(abs(v))

    clean: List[float] = []
    for v in values:
        vv = round(abs(float(v)), 6)
        if not (10 <= vv <= 1000):
            continue
        if vv in false_nums and not has_reliable_voltage_context(s):
            continue
        clean.append(vv)

    return sorted(set(clean))


def infer_voltage_polarity(*texts: Any) -> Optional[str]:
    joined = " ".join(str(t or "") for t in texts)
    if re.search(r"\+\s*\d+(?:\.\d+)?\s*(?:mV)?", joined):
        return "+"
    if re.search(r"-\s*\d+(?:\.\d+)?\s*(?:mV)?", joined):
        return "-"
    return None


def format_voltage_raw(v: float, polarity: Optional[str] = None) -> str:
    sign = polarity if polarity in {"+", "-"} else ""
    if float(v).is_integer():
        return f"{sign}{int(v)} mV"
    return f"{sign}{v:g} mV"


def extract_voltage_series_from_record(r: Dict[str, Any]) -> Tuple[List[float], Optional[str], Optional[str]]:
    polarity = clean_unknown(r.get("voltage_polarity"))
    if polarity not in {"+", "-"}:
        polarity = infer_voltage_polarity(
            r.get("voltage_raw"),
            r.get("voltage_series_raw"),
            r.get("caption_text"),
            r.get("evidence"),
            r.get("in_text_reference_evidence"),
        )

    false_nums = extract_numbers_from_mutation_or_peptide_names(
        r.get("analyte_name"),
        r.get("analytes"),
        r.get("analyte_alias"),
        r.get("nanopore_variant"),
        r.get("nanopore_name"),
        r.get("caption_text"),
        r.get("evidence"),
        r.get("in_text_reference_evidence"),
    )

    notes: List[str] = []
    series: List[float] = []

    # LLM 明确给的列表
    raw_series = (
        r.get("voltage_series_mV")
        or r.get("voltage_values_mV")
        or r.get("voltages_mV")
        or r.get("voltage_list_mV")
    )
    if isinstance(raw_series, list):
        for x in raw_series:
            v = num_to_float(x)
            if v is None:
                continue
            vv = round(abs(float(v)), 6)
            if vv in false_nums:
                notes.append(f"excluded {vv:g} from voltage_series_mV because it matches peptide/mutation numbering")
                continue
            series.append(vv)
    elif raw_series is not None:
        series.extend(parse_voltage_series_from_text(raw_series, false_nums=false_nums, require_context=False))

    joined = " | ".join(
        str(x or "") for x in [
            r.get("voltage_series_raw"),
            r.get("voltage_raw"),
            r.get("x_axis_label"),
            r.get("caption_text"),
            r.get("in_text_reference_evidence"),
            r.get("evidence"),
            r.get("ambiguous_reason"),
        ]
    )
    parsed = parse_voltage_series_from_text(joined, false_nums=false_nums, require_context=True)
    if parsed:
        series.extend(parsed)
        if len(parsed) > 1:
            notes.append("voltage series parsed from reliable voltage axis/caption/text context")

    # 单个 voltage_mV 要有可靠上下文且不能是名称编号
    v_single = num_to_float(r.get("voltage_mV"))
    single_context = " | ".join(
        str(x or "") for x in [
            r.get("voltage_raw"),
            r.get("x_axis_label"),
            r.get("caption_text"),
            r.get("in_text_reference_evidence"),
            r.get("evidence"),
        ]
    )
    if v_single is not None:
        vv = round(abs(float(v_single)), 6)
        if vv < 10:
            notes.append(f"excluded {vv:g} mV candidate because it is below minimum voltage threshold (10 mV)")
        elif vv in false_nums:
            notes.append(f"excluded {vv:g} mV candidate because it matches analyte/mutation numbering")
        elif has_reliable_voltage_context(single_context):
            series.append(vv)
        else:
            notes.append(f"excluded {vv:g} mV candidate because no reliable voltage context was found")

    series = sorted(set(round(float(v), 6) for v in series if v is not None and 10 <= abs(float(v)) <= 1000))

    if false_nums:
        filtered = []
        for v in series:
            if round(abs(float(v)), 6) in false_nums:
                notes.append(f"removed {v:g} because it matches peptide/mutation numbering")
                continue
            filtered.append(v)
        series = sorted(set(filtered))

    return series, polarity, "; ".join(dict.fromkeys(notes)) or None


def extract_concentration(raw: Any) -> Tuple[Optional[float], Optional[str], Optional[str]]:
    if raw is None:
        return None, None, None
    s = str(raw).strip()
    if not s:
        return None, None, None
    m = re.search(
        r"([-+]?\d+(?:\.\d+)?)\s*(μM|uM|µM|microM|mM|nM|pM|M|"
        r"mg\s*/\s*mL|ng\s*/\s*mL|μg\s*/\s*mL|ug\s*/\s*mL)",
        s,
        flags=re.I,
    )
    if not m:
        return None, None, s
    val = safe_float(m.group(1))
    unit = m.group(2)
    unit = unit.replace("µ", "μ")
    if unit.lower() == "um":
        unit = "μM"
    unit = re.sub(r"\s+", "", unit)
    return val, unit, s


def clean_electrolyte(val: Any) -> Optional[str]:
    """只保留电解质中的主盐成分，去除 EDTA 等添加剂。

    "1 M KCl, 1 mM EDTA" -> "1 M KCl"
    """
    val = clean_unknown(val)
    if not val:
        return None
    # 取第一个逗号或分号之前的部分
    first = re.split(r"[,;]", val, maxsplit=1)[0].strip()
    return first or None


def metric_category_from_text(*texts: Any) -> Optional[str]:
    joined = " | ".join(str(x or "") for x in texts)
    if NOISE_RE.search(joined) and not (TARGET_CURRENT_RE.search(joined) or TARGET_TIME_RE.search(joined)):
        return None
    if TARGET_TIME_RE.search(joined):
        return "dwell_time"
    if TARGET_CURRENT_RE.search(joined):
        # 更细分：Ib/I0/residual 是 residual_current；pA/ΔI/current drop 是 blockade_current/current_drop
        if re.search(r"I\s*_?b\s*/\s*I\s*_?0|I\s*/\s*I\s*0|residual|relative|normalized|残余|相对|归一", joined, re.I):
            return "residual_current"
        return "blockade_current"
    return None


def is_target_record(r: Dict[str, Any]) -> bool:
    own_axis_metric = " | ".join(
        str(r.get(k) or "") for k in [
            "metric", "metric_category",
            "dwell_current_metric", "dwell_time_metric",
            "x_axis_label", "y_axis_label",
            "x_axis_metric", "y_axis_metric", "signal_unit"
        ]
    )

    # 当前记录自身是 σb-only 就不输出
    if NOISE_RE.search(own_axis_metric) and not (TARGET_CURRENT_RE.search(own_axis_metric) or TARGET_TIME_RE.search(own_axis_metric)):
        return False

    # 检查新字段或旧字段是否有目标指标
    has_current_metric = bool(clean_unknown(r.get("dwell_current_metric")))
    has_time_metric = bool(clean_unknown(r.get("dwell_time_metric")))

    if not has_current_metric and not has_time_metric:
        # 回退到旧字段检查
        if not metric_category_from_text(
            r.get("metric"),
            r.get("metric_category"),
            r.get("x_axis_label"),
            r.get("y_axis_label"),
            r.get("blockade_current"),
            r.get("blockade_time"),
            r.get("residual_current"),
        ):
            return False

    # 只要满足待测物、纳米孔、图注、横纵坐标等基本上下文条件，即使有多个 panel 也全部保留。
    # 图表类型的筛选（如一维/二维/三维）将在后续步骤中由其他文件处理。
    return True

def has_sufficient_context(r: Dict[str, Any]) -> bool:
    """检查该图对应的完整上下文中是否包含足够的关键信息。

    上下文来源优先级：
    1. 图注（caption_text）
    2. 正文中引用该图的位置（in_text_reference_evidence）
    3. 图附近上下文（evidence）
    4. 方法部分 / 全局实验条件（通过 info_source_notes 体现）

    只要能从任意来源可靠证明关键信息即可，不要求所有信息都出现在图注中。
    至少需要确认：待测物、纳米孔、信号指标中的至少两项。
    """
    caption = str(r.get("caption_text") or "")
    evidence = str(r.get("evidence") or "")
    in_text = str(r.get("in_text_reference_evidence") or "")
    context = f"{caption} {evidence} {in_text}"

    checks = [
        # 待测物（可从图注、正文引用、上下文中确认）
        bool(clean_unknown(r.get("analyte_name")) or clean_unknown(r.get("analytes"))
             or re.search(r"peptide|protein|compound|analyte|Aβ|Abeta|amyloid", context, re.I)),
        # 纳米孔（可从图注、正文引用、上下文中确认）
        bool(clean_unknown(r.get("nanopore_name"))
             or re.search(r"nanopore|aerolysin|hemolysin|α-hl|MspA|ClyA|OmpF|OmpG|solid.state", context, re.I)),
        # 信号指标（可从图注、正文引用、上下文中确认）
        bool(TARGET_CURRENT_RE.search(context) or TARGET_TIME_RE.search(context)
             or clean_unknown(r.get("metric"))
             or clean_unknown(r.get("dwell_current_metric"))
             or clean_unknown(r.get("dwell_time_metric"))),
    ]
    # 至少需要确认待测物、纳米孔、信号指标中的至少两项
    return sum(checks) >= 2

def axis_role_from_record(r: Dict[str, Any]) -> Optional[str]:
    x_text = " | ".join(str(r.get(k) or "") for k in ["x_axis_label", "x_axis_metric"])
    y_text = " | ".join(str(r.get(k) or "") for k in ["y_axis_label", "y_axis_metric"])

    if (TARGET_CURRENT_RE.search(x_text) or TARGET_TIME_RE.search(x_text)) and COUNTS_RE.search(y_text):
        return "target_on_x_axis_counts_on_y_axis"
    if TARGET_CURRENT_RE.search(y_text) or TARGET_TIME_RE.search(y_text):
        return "target_on_y_axis"
    if TARGET_CURRENT_RE.search(x_text) or TARGET_TIME_RE.search(x_text):
        return "target_on_x_axis"
    return clean_unknown(r.get("axis_role"))


def analytes_from_record(r: Dict[str, Any]) -> List[str]:
    vals: List[str] = []
    if isinstance(r.get("analytes"), list):
        vals.extend(str(x).strip() for x in r.get("analytes") if str(x).strip())

    for k in ["analyte_name", "analyte"]:
        v = clean_unknown(r.get(k))
        if isinstance(v, str):
            # 兼容 "WT and A21G Aβ18-26" 这种不拆分不完全的情况；
            # 最好让 LLM 输出 analytes 列表，这里只做轻量兜底。
            vals.append(v)

    # 如果写成 "WT and A21G Aβ18-26 peptides"，不强行复杂拆分，避免误拆；
    # 但常见 Figure S6 可以做一个简单修正。
    expanded: List[str] = []
    for v in vals:
        s = str(v).strip()
        m = re.match(r"^(WT)\s+and\s+(A\d+[A-Z])\s+(Aβ\s*\d+\s*[-–]\s*\d+|Aβ\d+\s*[-–]\s*\d+|Abeta\s*\d+\s*[-–]\s*\d+)", s, flags=re.I)
        if m:
            base = re.sub(r"\s+", "", m.group(3)).replace("Aβ", "Aβ")
            expanded.append(f"{m.group(1)} {base}")
            expanded.append(f"{m.group(2)} {base}")
        else:
            expanded.append(s)

    return [str(x) for x in deduplicate_list(expanded)]


def value_map_by_voltage(r: Dict[str, Any]) -> Dict[float, Dict[str, Any]]:
    raw = (
        r.get("values_by_voltage")
        or r.get("signal_values_by_voltage")
        or r.get("metric_values_by_voltage")
        or r.get("y_values_by_voltage")
    )
    out: Dict[float, Dict[str, Any]] = {}
    if raw is None:
        return out

    if isinstance(raw, str):
        try:
            raw = json.loads(raw)
        except Exception:
            return out

    if isinstance(raw, dict):
        for k, v in raw.items():
            voltage = num_to_float(k)
            if voltage is None:
                continue
            if isinstance(v, dict):
                item = dict(v)
                item.setdefault("value", v.get("value") or v.get("y") or v.get("mean"))
            else:
                item = {"value": v}
            out[abs(float(voltage))] = item
        return out

    if isinstance(raw, list):
        for item in raw:
            if not isinstance(item, dict):
                continue
            voltage = num_to_float(item.get("voltage_mV") or item.get("voltage") or item.get("x"))
            if voltage is None:
                continue
            out[abs(float(voltage))] = item

    return out


def value_for_voltage(values_map: Dict[float, Dict[str, Any]], voltage: Any) -> Dict[str, Any]:
    if voltage is None:
        return {}
    vv = abs(float(voltage))
    for k, item in values_map.items():
        if abs(float(k) - vv) < 1e-6:
            return item
    for k, item in values_map.items():
        if round(float(k), 3) == round(vv, 3):
            return item
    return {}




def _normalize_single_record(
    r: Dict[str, Any],
    i: int,
    rows: List[Dict[str, Any]],
    folder_name: str,
    title: Optional[str],
    doi: Optional[str],
) -> None:
    analytes = analytes_from_record(r)
    if not analytes:
        analytes = [None]  # type: ignore[list-item]

    series, polarity, voltage_note = extract_voltage_series_from_record(r)
    if not series:
        series = [None]  # type: ignore[list-item]

    # 读取 panel 映射：哪个 panel 对应电流，哪个对应时间
    panel_dwell_current = clean_unknown(r.get("panel_dwell_current"))
    panel_dwell_time = clean_unknown(r.get("panel_dwell_time"))

    # 读取各类指标的 metric 名称
    dwell_current_metric = clean_unknown(r.get("dwell_current_metric"))
    dwell_time_metric = clean_unknown(r.get("dwell_time_metric"))

    # 如果 LLM 没有给出新字段，尝试从旧字段兼容
    if not panel_dwell_current and not panel_dwell_time:
        old_panel = clean_unknown(r.get("panel") or r.get("figure_panel"))
        old_metric_category = metric_category_from_text(
            r.get("metric"), r.get("metric_category"),
            r.get("x_axis_label"), r.get("y_axis_label"),
            r.get("blockade_current"), r.get("blockade_time"), r.get("residual_current"),
        )
        if old_panel:
            if old_metric_category in ["residual_current", "blockade_current"]:
                panel_dwell_current = old_panel
            elif old_metric_category == "dwell_time":
                panel_dwell_time = old_panel

    # 回退：从旧 metric 字段读取
    if not dwell_current_metric and not dwell_time_metric:
        old_metric = clean_unknown(r.get("metric"))
        old_metric_category = metric_category_from_text(
            r.get("metric"), r.get("metric_category"),
            r.get("x_axis_label"), r.get("y_axis_label"),
            r.get("blockade_current"), r.get("blockade_time"), r.get("residual_current"),
        )
        if old_metric_category in ["residual_current", "blockade_current"]:
            dwell_current_metric = old_metric
        elif old_metric_category == "dwell_time":
            dwell_time_metric = old_metric

    conc_raw = (
        r.get("analyte_concentration_raw")
        or r.get("analyte_concentration")
        or r.get("final_concentration")
        or r.get("sample_concentration")
        or r.get("concentration")
    )
    conc_value, conc_unit, conc_raw_norm = extract_concentration(conc_raw)
    if r.get("analyte_concentration_value") is not None:
        conc_value = safe_float(r.get("analyte_concentration_value"))
    if clean_unknown(r.get("analyte_concentration_unit")):
        conc_unit = clean_unknown(r.get("analyte_concentration_unit"))

    # 分别读取电流和时间的 values_by_voltage
    current_values_map = value_map_by_voltage({
        "values_by_voltage": r.get("dwell_current_values_by_voltage"),
    }) if r.get("dwell_current_values_by_voltage") else value_map_by_voltage(r)

    time_values_map = value_map_by_voltage({
        "values_by_voltage": r.get("dwell_time_values_by_voltage"),
    }) if r.get("dwell_time_values_by_voltage") else {}

    # 如果 LLM 只给了旧格式的 values_by_voltage，根据 metric_category 分配
    if not r.get("dwell_current_values_by_voltage") and not r.get("dwell_time_values_by_voltage"):
        old_metric_category = metric_category_from_text(
            r.get("metric"), r.get("metric_category"),
            r.get("x_axis_label"), r.get("y_axis_label"),
            r.get("blockade_current"), r.get("blockade_time"), r.get("residual_current"),
        )
        if old_metric_category == "dwell_time":
            time_values_map = current_values_map
            current_values_map = {}

    base_record_id = clean_unknown(r.get("record_id")) or f"R{i}"

    for analyte_idx, analyte in enumerate(analytes, start=1):
        for v_idx, voltage in enumerate(series, start=1):
            if voltage is None:
                voltage_raw = clean_unknown(r.get("voltage_raw"))
                voltage_mV = safe_float(r.get("voltage_mV"))
                if voltage_mV is not None and abs(float(voltage_mV)) < 10:
                    voltage_mV = None
                current_item: Dict[str, Any] = {}
                time_item: Dict[str, Any] = {}
            else:
                voltage_raw = format_voltage_raw(float(voltage), polarity)
                voltage_mV = int(voltage) if float(voltage).is_integer() else float(voltage)
                current_item = value_for_voltage(current_values_map, voltage)
                time_item = value_for_voltage(time_values_map, voltage)

            # 电流值
            dwell_current_val = safe_float(
                current_item.get("value") or current_item.get("y") or current_item.get("mean")
                or r.get("dwell_current_signal_value") or r.get("signal_value")
            )
            dwell_current_unit = clean_unknown(
                current_item.get("unit") or current_item.get("y_unit")
                or r.get("dwell_current_signal_unit") or r.get("signal_unit")
                or r.get("y_axis_unit")
            )

            # 时间值
            dwell_time_val = safe_float(
                time_item.get("value") or time_item.get("y") or time_item.get("mean")
                or r.get("dwell_time_signal_value")
            )
            dwell_time_unit = clean_unknown(
                time_item.get("unit") or time_item.get("y_unit")
                or r.get("dwell_time_signal_unit")
            )

            rec_id = str(base_record_id)
            if len(analytes) > 1:
                rec_id += f"_A{analyte_idx}"
            if len(series) > 1:
                rec_id += f"_V{v_idx}"

            # figure_id 拆分为 dwell_current 和 dwell_time 两部分
            figure_id_dwell_current = clean_unknown(r.get("figure_id_dwell_current")) or clean_unknown(r.get("figure_id") or r.get("figure_or_table"))
            figure_id_dwell_time = clean_unknown(r.get("figure_id_dwell_time")) or clean_unknown(r.get("figure_id") or r.get("figure_or_table"))
            # 兼容旧字段
            figure_id = clean_unknown(r.get("figure_id") or r.get("figure_or_table"))
            figure_group_id = clean_unknown(r.get("figure_group_id")) or figure_id
            condition_group_id = clean_unknown(r.get("condition_group_id")) or " | ".join(
                str(x or "") for x in [
                    figure_group_id,
                    r.get("nanopore_name"),
                    r.get("nanopore_variant"),
                    r.get("electrolyte"),
                    r.get("buffer"),
                    r.get("pH"),
                    conc_raw_norm,
                ]
                if clean_unknown(x)
            )

            rows.append({
                "folder": folder_name,
                "title": title,
                "doi": doi,

                "record_id": rec_id,
                "figure_id": figure_id,
                "figure_id_dwell_current": figure_id_dwell_current,
                "figure_id_dwell_time": figure_id_dwell_time,
                "figure_group_id": figure_group_id,
                "panel_dwell_current": panel_dwell_current,
                "panel_dwell_time": panel_dwell_time,
                "condition_group_id": condition_group_id,

                "analyte": clean_unknown(analyte),
                "analyte_alias": json_or_none(r.get("analyte_alias")),
                "analyte_type": clean_unknown(r.get("analyte_type")),

                "nanopore": clean_unknown(r.get("nanopore_name")),
                "nanopore_variant": clean_unknown(r.get("nanopore_variant")),
                "nanopore_type": clean_unknown(r.get("nanopore_type")),

                "dwell_current_metric": dwell_current_metric,
                "dwell_time_metric": dwell_time_metric,
                "chart_type": clean_unknown(r.get("chart_type")),
                "axis_role": axis_role_from_record(r),

                "voltage_raw": voltage_raw,
                "voltage_mV": voltage_mV,
                "voltage_polarity": polarity,
                "voltage_series_raw": clean_unknown(r.get("voltage_series_raw") or r.get("voltage_raw")),
                "voltage_series_mV_json": json.dumps(
                    [int(x) if x is not None and float(x).is_integer() else x for x in series if x is not None],
                    ensure_ascii=False,
                ),
                "voltage_inference_note": voltage_note,

                "dwell_current": dwell_current_val,
                "dwell_current_unit": dwell_current_unit,
                "dwell_time": dwell_time_val,
                "dwell_time_unit": dwell_time_unit,

                "analyte_concentration_raw": conc_raw_norm,
                "analyte_concentration_value": conc_value,
                "analyte_concentration_unit": conc_unit,
                "concentration_scope": clean_unknown(r.get("concentration_scope")),

                "chamber": clean_unknown(r.get("chamber")),
                "cis_trans_addition": clean_unknown(r.get("cis_trans_addition")),
                "electrolyte": clean_electrolyte(r.get("electrolyte")),
                "salt_concentration": clean_unknown(r.get("salt_concentration")),
                "buffer": clean_unknown(r.get("buffer")),
                "pH": safe_float(r.get("pH")),
                "temperature": clean_unknown(r.get("temperature")),

                "source_file": clean_unknown(r.get("source_file")),
                "page_range": clean_unknown(r.get("page_range")),
                "source_kind": clean_unknown(r.get("source_kind")),

                "caption_text": clean_unknown(r.get("caption_text") or r.get("figure_caption_evidence")),
                "in_text_reference_evidence": clean_unknown(r.get("in_text_reference_evidence") or r.get("text_reference_evidence")),
                "evidence": clean_unknown(r.get("evidence")),
                "paired_panels": json_or_none(r.get("paired_panels")),
                "panel_relation": clean_unknown(r.get("panel_relation")),
                "info_source_notes": json_or_none(r.get("info_source_notes")),
                "ambiguous_reason": clean_unknown(r.get("ambiguous_reason")),
                "confidence": safe_float(r.get("confidence")),
                "excluded_non_target_note": "σb/current fluctuation/standard deviation are not target outputs; they are kept only as contextual evidence if mentioned.",
            })


def normalize_records(
    records: List[Dict[str, Any]],
    folder_name: str,
    title: Optional[str],
    doi: Optional[str],
) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []

    for i, r in enumerate(records, start=1):
        if not isinstance(r, dict):
            continue
        if not is_target_record(r):
            continue
        if not has_sufficient_context(r):
            continue
        _normalize_single_record(r, i, rows, folder_name, title, doi)

    df = pd.DataFrame(rows)
    if df.empty:
        return df

    df = deduplicate_relation_rows(df)
    return df


def deduplicate_relation_rows(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df

    key_cols = [
        "figure_id_dwell_current", "figure_id_dwell_time", "panel_dwell_current", "panel_dwell_time",
        "analyte", "nanopore", "nanopore_variant",
        "dwell_current_metric", "dwell_time_metric", "voltage_mV", "pH", "electrolyte",
        "buffer", "analyte_concentration_raw"
    ]
    cols = [c for c in key_cols if c in df.columns]
    if not cols:
        return df.drop_duplicates()

    return df.drop_duplicates(subset=cols, keep="first").reset_index(drop=True)


# =============================================================================
# 12. Figure 条件汇总
# =============================================================================

def build_figure_condition_summary(signal_df: pd.DataFrame) -> pd.DataFrame:
    if signal_df.empty:
        return pd.DataFrame([{"status": "no_signal_relation_records"}])

    group_cols = [
        "folder", "figure_group_id", "figure_id", "figure_id_dwell_current", "figure_id_dwell_time",
        "panel_dwell_current", "panel_dwell_time",
        "nanopore", "nanopore_variant", "dwell_current_metric", "dwell_time_metric",
        "analyte_concentration_raw", "electrolyte", "buffer", "pH",
        "chamber", "cis_trans_addition", "temperature"
    ]
    group_cols = [c for c in group_cols if c in signal_df.columns]

    rows: List[Dict[str, Any]] = []
    for keys, g in signal_df.groupby(group_cols, dropna=False):
        if not isinstance(keys, tuple):
            keys = (keys,)
        row = {col: val for col, val in zip(group_cols, keys)}

        row["analyte_list"] = "; ".join(
            str(x) for x in deduplicate_list(g.get("analyte", pd.Series(dtype=str)).tolist())
        )
        row["voltage_series_mV"] = "; ".join(
            str(int(v) if isinstance(v, (int, float)) and float(v).is_integer() else v)
            for v in deduplicate_list(g.get("voltage_mV", pd.Series(dtype=float)).tolist())
        )
        row["source_files"] = "; ".join(
            str(x) for x in deduplicate_list(g.get("source_file", pd.Series(dtype=str)).tolist())
        )
        row["evidence_examples"] = " || ".join(
            str(x) for x in deduplicate_list(
                g.get("caption_text", pd.Series(dtype=str)).dropna().tolist()
                + g.get("evidence", pd.Series(dtype=str)).dropna().tolist()
            )[:3]
        )
        rows.append(row)

    return pd.DataFrame(rows)


# =============================================================================
# 13. 文件夹处理和 Excel 输出
# =============================================================================

def should_skip_folder(folder: Path) -> Tuple[bool, str]:
    name = folder.name.strip()
    if not name:
        return True, "empty_name"
    lower = name.lower()
    if lower in SKIP_FOLDERS:
        return True, "system_folder"
    if name.startswith(".") or (name.startswith("__") and name.endswith("__")):
        return True, "hidden_or_system_folder"
    return False, ""


def folder_has_supported_files(folder: Path) -> bool:
    try:
        return len(gather_supported_files(folder)) > 0
    except Exception:
        return False


def process_folder(
    folder: Path,
    llm: DoubaoArkLLM,
    use_vlm_pages: bool = False,
    vlm: Optional[DoubaoArkLLM] = None,
    vlm_max_pages_per_folder: int = 10,
) -> Tuple[pd.DataFrame, pd.DataFrame, Dict[str, Any]]:
    files = gather_supported_files(folder)
    if not files:
        return (
            pd.DataFrame(),
            pd.DataFrame([{"status": "no_supported_files"}]),
            {"error": "no_supported_files"},
        )

    print(f"处理文件夹: {folder}")
    print(f"识别到文件数量: {len(files)}")

    chunks: List[TextChunk] = []
    pages_map: Dict[str, List[Tuple[int, str]]] = {}
    all_texts: List[str] = []
    first_pages: List[Tuple[int, str]] = []
    file_roles: Dict[str, str] = {}

    for fp in files:
        role = classify_file_role(fp)
        file_roles[fp.name] = role
        try:
            pages = read_file_pages(fp)
            pages_map[fp.name] = pages
            if pages and not first_pages:
                first_pages = pages[:1]
            all_texts.extend(t for _, t in pages)

            cs = chunk_pages(fp.name, role, pages)
            chunks.extend(cs)
            print(f"  - {fp.name}: 提取到 {len(pages)} 页/段，切成 {len(cs)} 个 chunk")
        except Exception as exc:
            print(f"  - {fp.name}: 读取失败 -> {exc}")

    if not chunks:
        return (
            pd.DataFrame(),
            pd.DataFrame([{"status": "no_text_extracted"}]),
            {"error": "no_text_extracted", "file_roles": file_roles},
        )

    signal_chunks = select_signal_candidate_chunks(chunks, top_k=18, min_score=12.0)
    print(f"  阻断信号候选 chunks: {len(signal_chunks)}")

    vlm_chunks: List[TextChunk] = []
    if use_vlm_pages and vlm is not None:
        vlm_chunks = extract_vlm_chunks(
            files=files,
            pages_map=pages_map,
            signal_chunks=signal_chunks,
            vlm=vlm,
            max_pages=vlm_max_pages_per_folder,
        )
        print(f"  VLM 有效图表 chunks: {len(vlm_chunks)}")

    final_chunks = sorted(signal_chunks + vlm_chunks, key=lambda c: c.score, reverse=True)

    llm_result = llm_extract_signal_relations(llm, folder.name, final_chunks)
    title = coerce_title(llm_result.get("title"), first_pages)
    doi = coerce_doi(llm_result.get("doi"), all_texts)

    raw_records = llm_result.get("records") if isinstance(llm_result.get("records"), list) else []
    signal_df = normalize_records(raw_records, folder.name, title, doi)
    condition_df = build_figure_condition_summary(signal_df)

    print(f"  已抽取 阻断信号-待测物-纳米孔-电压/条件 关系: {len(signal_df)} 条")

    meta = {
        "title": title,
        "doi": doi,
        "file_roles": file_roles,
        "all_chunk_count": len(chunks),
        "signal_candidate_chunk_count": len(signal_chunks),
        "vlm_chunk_count": len(vlm_chunks),
        "raw_record_count_from_llm": len(raw_records),
        "normalized_signal_record_count": len(signal_df),
        "llm_raw_response": llm_result.get("_raw_response"),
        "llm_json": {k: v for k, v in llm_result.items() if k != "_raw_response"},
    }

    return signal_df, condition_df, meta


def autosize_excel_columns(writer: pd.ExcelWriter, sheet_name: str, df: pd.DataFrame) -> None:
    try:
        ws = writer.sheets[sheet_name]
        for idx, col in enumerate(df.columns, start=1):
            max_len = len(str(col))
            sample = df[col].astype(str).head(200).tolist()
            for val in sample:
                max_len = max(max_len, min(len(str(val)), 80))
            ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = min(max_len + 2, 60)
    except Exception:
        pass


def save_folder_excel(
    folder: Path,
    signal_df: pd.DataFrame,
    condition_df: pd.DataFrame,
    meta: Dict[str, Any],
) -> Path:
    out = folder / f"{folder.name}_nanopore_blockade_signal_relations.xlsx"

    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        if signal_df.empty:
            df1 = pd.DataFrame([{"status": "no_signal_relation_records"}])
        else:
            df1 = _filter_and_enhance_columns(signal_df)
        df1.to_excel(writer, index=False, sheet_name="signal_relation_records")
        autosize_excel_columns(writer, "signal_relation_records", df1)

        df2 = condition_df if not condition_df.empty else pd.DataFrame([{"status": "no_figure_condition_summary"}])
        df2.to_excel(writer, index=False, sheet_name="figure_condition_summary")
        autosize_excel_columns(writer, "figure_condition_summary", df2)

        meta_rows = []
        for k, v in meta.items():
            if isinstance(v, (dict, list)):
                v = json.dumps(v, ensure_ascii=False, indent=2)
            meta_rows.append({"key": k, "value": v})
        df_meta = pd.DataFrame(meta_rows)
        df_meta.to_excel(writer, index=False, sheet_name="meta")
        autosize_excel_columns(writer, "meta", df_meta)

    return out


def _filter_and_enhance_columns(df: pd.DataFrame) -> pd.DataFrame:
    """筛选列：保留用户需要的列。"""
    # 需要保留的列（dwell_time 和 dwell_current 已在提取时根据 metric 确定）
    keep_cols = [
        "folder", "title", "doi",
        "analyte", "nanopore", "nanopore_variant",
        "salt_concentration", "electrolyte", "pH",
        "analyte_concentration_value", "analyte_concentration_unit",
        "voltage_raw",
        "figure_id_dwell_current", "panel_dwell_current", "dwell_current_metric", "dwell_current",
        "figure_id_dwell_time", "panel_dwell_time", "dwell_time_metric", "dwell_time",
    ]

    # 确保所有需要的列都存在，不存在则填 None
    for col in keep_cols:
        if col not in df.columns:
            df[col] = None

    return df[keep_cols].copy()


def _is_current_metric(metric: Any) -> bool:
    """判断是否为阻断电流类指标（Ib/I0 等）。"""
    if not metric:
        return False
    m = str(metric).lower()
    return any(kw in m for kw in [
        "ib/i0", "i/i0", "residual_current", "blockade_current",
        "current_drop", "current_blockage", "relative_current",
        "normalized_current", "i_b/i_0"
    ])


def _is_time_metric(metric: Any) -> bool:
    """判断是否为阻断时间类指标（tD/log(tD) 等）。"""
    if not metric:
        return False
    m = str(metric).lower()
    return any(kw in m for kw in [
        "td", "t_d", "log(td", "log(t_d", "dwell_time",
        "blockade_time", "blockage_time", "residence_time",
        "translocation_time", "event_duration"
    ])


def process_root(
    root_dir: Path,
    model: Optional[str] = None,
    base_url: Optional[str] = None,
    output_name: str = "batch_nanopore_blockade_signal_summary.xlsx",
    use_vlm_pages: bool = False,
    vlm_model: Optional[str] = None,
    vlm_max_pages_per_folder: int = 10,
) -> None:
    if not root_dir.exists() or not root_dir.is_dir():
        raise NotADirectoryError(f"根目录不存在或不是文件夹：{root_dir}")

    llm = DoubaoArkLLM(model=model, base_url=base_url)
    vlm = (
        DoubaoArkLLM(
            model=vlm_model or os.getenv("DOUBAO_VL_MODEL") or os.getenv("ARK_VL_MODEL") or os.getenv("DOUBAO_MODEL") or os.getenv("ARK_MODEL") or "doubao-seed-2-0-pro-260215",
            base_url=base_url,
            timeout=180,
            max_retries=4,
        )
        if use_vlm_pages else None
    )

    raw_subfolders = sorted([p for p in root_dir.iterdir() if p.is_dir()])
    if not raw_subfolders:
        raw_subfolders = [root_dir]

    subfolders: List[Path] = []
    for f in raw_subfolders:
        skip, reason = should_skip_folder(f)
        if skip:
            print("#" * 100)
            print(f"跳过文件夹: {f.name}（{reason}）")
            continue

        if f != root_dir and not folder_has_supported_files(f):
            print("#" * 100)
            print(f"跳过文件夹: {f.name}（no_supported_files）")
            continue

        subfolders.append(f)

    if not subfolders and folder_has_supported_files(root_dir):
        subfolders = [root_dir]

    summary_rows: List[Dict[str, Any]] = []

    for idx, folder in enumerate(subfolders, start=1):
        print("#" * 100)
        print(f"[{idx}/{len(subfolders)}] {folder.name}")

        try:
            signal_df, condition_df, meta = process_folder(
                folder=folder,
                llm=llm,
                use_vlm_pages=use_vlm_pages,
                vlm=vlm,
                vlm_max_pages_per_folder=vlm_max_pages_per_folder,
            )
            out = save_folder_excel(folder, signal_df, condition_df, meta)

            summary_rows.append({
                "folder": folder.name,
                "output_excel": str(out),
                "signal_relation_count": len(signal_df),
                "figure_condition_count": 0 if condition_df.empty else len(condition_df),
                "title": meta.get("title"),
                "doi": meta.get("doi"),
                "status": "ok",
            })
            print(f"  Excel 已保存: {out}")

        except Exception as exc:
            traceback.print_exc()
            summary_rows.append({
                "folder": folder.name,
                "output_excel": None,
                "signal_relation_count": 0,
                "figure_condition_count": 0,
                "title": None,
                "doi": None,
                "status": f"failed: {exc}",
            })
            print(f"  -> 处理失败: {exc}")

    if summary_rows:
        summary_path = root_dir / output_name
        df_sum = pd.DataFrame(summary_rows)
        with pd.ExcelWriter(summary_path, engine="openpyxl") as writer:
            df_sum.to_excel(writer, index=False, sheet_name="summary")
            autosize_excel_columns(writer, "summary", df_sum)
        print("=" * 100)
        print(f"批处理汇总已保存: {summary_path}")


# =============================================================================
# 14. main
# =============================================================================

def main() -> None:
    parser = argparse.ArgumentParser(
        description="批量抽取纳米孔阻断电流/阻断时间与待测物、纳米孔、电压、实验条件和浓度的对应关系"
    )
    parser.add_argument(
        "root_positional",
        nargs="?",
        default=None,
        help="根目录路径，可省略；省略时默认使用脚本所在目录",
    )
    parser.add_argument(
        "--root",
        dest="root_named",
        default=None,
        help="根目录路径",
    )
    parser.add_argument(
        "--model",
        default=(os.getenv("DOUBAO_MODEL") or os.getenv("ARK_MODEL") or os.getenv("VOLCENGINE_MODEL") or "doubao-seed-2-0-pro-260215"),
        help="豆包/火山方舟文本模型名或推理接入点 ID，默认读取 DOUBAO_MODEL/ARK_MODEL/VOLCENGINE_MODEL",
    )
    parser.add_argument(
        "--base-url",
        default=(os.getenv("ARK_BASE_URL") or os.getenv("DOUBAO_BASE_URL") or DEFAULT_BASE_URL),
        help="豆包/火山方舟 OpenAI-compatible Chat Completions 接口地址",
    )
    parser.add_argument(
        "--output-name",
        default="batch_nanopore_blockade_signal_summary.xlsx",
        help="根目录汇总 Excel 文件名",
    )
    parser.add_argument(
        "--use-vlm-pages",
        action="store_true",
        help="z",
    )
    parser.add_argument(
        "--vlm-model",
        default=(os.getenv("DOUBAO_VL_MODEL") or os.getenv("ARK_VL_MODEL") or os.getenv("DOUBAO_MODEL") or os.getenv("ARK_MODEL") or "doubao-seed-2-0-pro-260215"),
        help="豆包/火山方舟视觉模型名或视觉推理接入点 ID，仅 --use-vlm-pages 时使用",
    )
    parser.add_argument(
        "--vlm-max-pages-per-folder",
        type=int,
        default=10,
        help="每个文件夹最多识别多少个候选 PDF 页面",
    )

    args = parser.parse_args()

    root_value = args.root_named or args.root_positional
    if root_value:
        root = Path(root_value)
    else:
        root = Path(__file__).resolve().parent
        print(f"未提供根目录，默认使用脚本所在目录: {root}")

    if not root.exists():
        print(f"根目录不存在: {root}")
        sys.exit(1)

    process_root(
        root_dir=root,
        model=args.model,
        base_url=args.base_url,
        output_name=args.output_name,
        use_vlm_pages=args.use_vlm_pages,
        vlm_model=args.vlm_model,
        vlm_max_pages_per_folder=args.vlm_max_pages_per_folder,
    )


if __name__ == "__main__":
    main()